"""CSV/Excel bank statement parsing — Phase 18-MVP-2."""

from __future__ import annotations

import csv
import datetime
import io
import json
import logging
import unicodedata
from html.parser import HTMLParser
from typing import Any

_log = logging.getLogger(__name__)

import pandas as pd

from reconciliation.amounts import parse_amount_str
from reconciliation.normalize import normalize_description

CANONICAL_FIELDS = (
    "date",
    "description",
    "amount",
    "debit",
    "credit",
    "balance",
    "bank_reference",
)

# Preset column maps: canonical_field -> list of header aliases (lowercase)
COLUMN_PRESETS: dict[str, dict[str, tuple[str, ...]]] = {
    "generic_en": {
        "date": ("date", "transaction date", "posting date", "value date"),
        "description": ("description", "details", "narrative", "memo"),
        "amount": ("amount", "transaction amount"),
        "debit": ("debit", "withdrawal", "out"),
        "credit": ("credit", "deposit", "in"),
        "balance": ("balance", "running balance"),
        "bank_reference": ("reference", "bank reference", "ref", "check no", "check number"),
    },
    "generic_tr": {
        "date": (
            "tarih",
            "işlem tarihi",
            "islem tarihi",
            "valör tarihi",
            "valor tarihi",
            "ekstre tarihi",
            "statement tarih",
            "statement tarihi",
        ),
        "description": (
            "açıklama",
            "aciklama",
            "işlem açıklaması",
            "islem aciklamasi",
            "detay",
            "açıklamalar",
        ),
        "amount": ("tutar", "işlem tutarı", "islem tutari", "meblağ", "meblag"),
        "debit": ("borç", "borc", "çıkış", "cikis", "borç tutarı", "borc tutari"),
        "credit": ("alacak", "giriş", "giris", "alacak tutarı", "alacak tutari"),
        "balance": ("bakiye", "güncel bakiye", "guncel bakiye", "kalan bakiye"),
        "bank_reference": (
            "referans",
            "dekont no",
            "fiş no",
            "fis no",
            "işlem no",
            "islem no",
            "sira no",
            "sıra no",
        ),
    },
}

# Substring keywords when exact alias match fails (Turkish bank exports).
_FIELD_KEYWORDS: dict[str, tuple[str, ...]] = {
    "date": ("tarih", "date", "valör", "valor"),
    "description": ("açıklama", "aciklama", "description", "detay", "explanation"),
    "amount": ("tutar", "amount", "meblağ", "meblag"),
    "debit": ("borç", "borc", "withdrawal", "debit"),
    "credit": ("alacak", "credit", "deposit"),
    "balance": ("bakiye", "balance"),
    "bank_reference": ("referans", "reference", "dekont", "fiş", "fis"),
}

_DATE_FORMATS = (
    "%Y-%m-%d",
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d-%m-%Y",
)


def _ext_from_name(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _decode_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _is_xlsx_bytes(file_bytes: bytes) -> bool:
    return len(file_bytes) >= 4 and file_bytes[:4] == b"PK\x03\x04"


def _is_ole_xls_bytes(file_bytes: bytes) -> bool:
    return len(file_bytes) >= 8 and file_bytes[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def detect_file_format(file_bytes: bytes, filename: str) -> str:
    """Detect actual file format (Turkish banks often export HTML/XML as .xls/.xlsx)."""
    if not file_bytes:
        return "empty"
    if _is_xlsx_bytes(file_bytes):
        return "xlsx"
    if _is_ole_xls_bytes(file_bytes):
        return "xls_ole"
    head_b = file_bytes[:4096].lstrip()
    head = _decode_text(file_bytes[:4096]).lower()
    if head_b.startswith(b"<?xml") or head.startswith("<?xml"):
        if "schemas-microsoft-com:office:spreadsheet" in head or "ss:workbook" in head:
            return "spreadsheetml"
    if (
        head_b[:1] == b"<"
        or "<table" in head
        or "<html" in head
        or "<meta" in head
        or "<head" in head
    ):
        return "html"
    if "," in head or ";" in head:
        return "csv"
    if _ext_from_name(filename) in ("xlsx", "xls"):
        return "excel_unrecognized"
    if _ext_from_name(filename) == "csv":
        return "csv"
    return "unknown"


def is_real_xlsx(file_bytes: bytes, filename: str) -> bool:
    return detect_file_format(file_bytes, filename) == "xlsx"


def read_tabular_preview(
    file_bytes: bytes,
    filename: str,
    *,
    header_row: int = 1,
    sheet_name: str | None = None,
    max_rows: int = 6,
) -> tuple[list[str], list[dict[str, Any]]]:
    """Return column headers and first data rows (for mapping UI)."""
    df = _read_dataframe(
        file_bytes, filename, header_row=header_row, sheet_name=sheet_name
    )
    headers = [str(c) for c in df.columns]
    preview = df.head(max_rows).fillna("").astype(str).to_dict(orient="records")
    return headers, preview


def _alias_matches_header(low: str, alias: str) -> bool:
    """Match header to alias without false positives (e.g. Tutar ≠ Borç tutarı)."""
    alias = _ascii_fold(alias).lower().strip()
    if not alias or not low:
        return False
    if low == alias:
        return True
    words = low.split()
    if alias in words:
        return True
    if low.endswith(" " + alias) or low.startswith(alias + " "):
        return True
    return False


def suggest_column_mapping(headers: list[str]) -> dict[str, str | None]:
    """Auto-suggest mapping from file headers to canonical fields."""
    lower_map = {h: _header_lower(h) for h in headers}
    mapping: dict[str, str | None] = {f: None for f in CANONICAL_FIELDS}
    used_headers: set[str] = set()
    # Prefer signed amount (Tutar) before separate Borç/Alacak columns.
    field_order = (
        "date",
        "description",
        "amount",
        "debit",
        "credit",
        "balance",
        "bank_reference",
    )
    for field in field_order:
        for preset in COLUMN_PRESETS.values():
            aliases = preset.get(field, ())
            for header, low in lower_map.items():
                if header in used_headers:
                    continue
                if any(_alias_matches_header(low, a) for a in aliases):
                    mapping[field] = header
                    used_headers.add(header)
                    break
            if mapping.get(field):
                break
        if mapping.get(field):
            continue
        keywords = _FIELD_KEYWORDS.get(field, ())
        for header, low in lower_map.items():
            if header in used_headers:
                continue
            if any(kw in low for kw in keywords):
                mapping[field] = header
                used_headers.add(header)
                break
    return mapping


def _ascii_fold(text: str) -> str:
    """Fold Turkish/diacritic chars for robust header matching (Borç → borc)."""
    s = unicodedata.normalize("NFKD", text)
    return "".join(c for c in s if not unicodedata.combining(c))


def _header_lower(header: str) -> str:
    """Normalize header text for matching (handles 'Statement Tarih', newlines, etc.)."""
    s = str(header).strip().replace("\n", " ").replace("\r", " ")
    s = " ".join(s.split())
    return _ascii_fold(s).lower()


def _apply_header_row(df: pd.DataFrame, header_row: int) -> pd.DataFrame:
    header_idx = max(0, header_row - 1)
    if header_idx >= len(df):
        return df
    out = df.copy()
    out.columns = [str(c).strip() for c in out.iloc[header_idx].tolist()]
    out = out.iloc[header_idx + 1 :].reset_index(drop=True)
    return out


class _HtmlTableExtractor(HTMLParser):
    """Minimal HTML table parser for Turkish bank .xls/.xlsx exports (stdlib only)."""

    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._in_table = False
        self._in_row = False
        self._in_cell = False
        self._current_table: list[list[str]] = []
        self._current_row: list[str] = []
        self._cell_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self._in_table = True
            self._current_table = []
        elif self._in_table and tag == "tr":
            self._in_row = True
            self._current_row = []
        elif self._in_row and tag in ("td", "th"):
            self._in_cell = True
            self._cell_parts = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._in_cell:
            self._current_row.append("".join(self._cell_parts).strip())
            self._in_cell = False
        elif tag == "tr" and self._in_row:
            if self._current_row:
                self._current_table.append(self._current_row)
            self._in_row = False
        elif tag == "table" and self._in_table:
            if self._current_table:
                self.tables.append(self._current_table)
            self._in_table = False

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._cell_parts.append(data)


def _html_tables_stdlib(text: str) -> list[pd.DataFrame]:
    parser = _HtmlTableExtractor()
    parser.feed(text)
    tables: list[pd.DataFrame] = []
    for raw in parser.tables:
        if not raw:
            continue
        width = max(len(r) for r in raw)
        rows = [r + [""] * (width - len(r)) for r in raw]
        tables.append(pd.DataFrame(rows))
    return tables


def _read_html_tables_raw(file_bytes: bytes) -> list[pd.DataFrame]:
    text = _decode_text(file_bytes)
    tables: list[pd.DataFrame] = []
    try:
        tables = pd.read_html(io.StringIO(text), header=None)
    except ImportError:
        tables = []
    except Exception:
        _log.debug("read_html failed, falling back to stdlib parser", exc_info=True)
        tables = []
    tables = [t for t in tables if len(t) > 0 and len(t.columns) > 0]
    if tables:
        return tables
    tables = _html_tables_stdlib(text)
    if not tables:
        raise ValueError("No table found in HTML bank export")
    return tables


def _pick_largest_table(tables: list[pd.DataFrame]) -> pd.DataFrame:
    if not tables:
        raise ValueError("No table found in HTML bank export")
    return max(tables, key=lambda t: len(t) * max(1, len(t.columns)))


def _read_spreadsheetml_raw(file_bytes: bytes, max_rows: int | None = None) -> pd.DataFrame:
    import xml.etree.ElementTree as ET

    root = ET.fromstring(_decode_text(file_bytes))
    rows_data: list[list[str]] = []
    for elem in root.iter():
        if not elem.tag.endswith("Row"):
            continue
        row_cells: list[str] = []
        col_idx = 0
        for cell in elem:
            if not cell.tag.endswith("Cell"):
                continue
            idx_attr = None
            for key, val in cell.attrib.items():
                if key.endswith("Index"):
                    idx_attr = int(val)
                    break
            if idx_attr is not None:
                while len(row_cells) < idx_attr - 1:
                    row_cells.append("")
                col_idx = idx_attr - 1
            data_el = next((c for c in cell if c.tag.endswith("Data")), None)
            val = (data_el.text or "").strip() if data_el is not None else ""
            while len(row_cells) < col_idx:
                row_cells.append("")
            if len(row_cells) == col_idx:
                row_cells.append(val)
            elif col_idx < len(row_cells):
                row_cells[col_idx] = val
            col_idx += 1
        if row_cells:
            rows_data.append(row_cells)
        if max_rows is not None and len(rows_data) >= max_rows:
            break
    if not rows_data:
        raise ValueError("No rows found in SpreadsheetML export")
    width = max(len(r) for r in rows_data)
    for row in rows_data:
        while len(row) < width:
            row.append("")
    return pd.DataFrame(rows_data)


def _scan_rows_for_header(
    file_bytes: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    max_scan: int = 15,
) -> list[list[str]]:
    fmt = detect_file_format(file_bytes, filename)
    scanned: list[list[str]] = []

    def _append_from_df(df: pd.DataFrame) -> None:
        for idx in range(min(max_scan, len(df))):
            cells = [
                str(v).strip()
                for v in df.iloc[idx].tolist()
                if v is not None and not (isinstance(v, float) and pd.isna(v)) and str(v).strip()
            ]
            scanned.append(cells)

    try:
        if fmt == "xlsx":
            raw = pd.read_excel(
                io.BytesIO(file_bytes),
                header=None,
                nrows=max_scan,
                engine="openpyxl",
                sheet_name=sheet_name if sheet_name else 0,
            )
            _append_from_df(raw)
        elif fmt == "html":
            _append_from_df(_pick_largest_table(_read_html_tables_raw(file_bytes)).head(max_scan))
        elif fmt == "spreadsheetml":
            _append_from_df(_read_spreadsheetml_raw(file_bytes, max_rows=max_scan))
        elif fmt == "excel_unrecognized":
            for reader in (_read_html_tables_raw, _read_spreadsheetml_raw):
                try:
                    if reader is _read_html_tables_raw:
                        df = _pick_largest_table(reader(file_bytes)).head(max_scan)
                    else:
                        df = reader(file_bytes, max_rows=max_scan)
                    _append_from_df(df)
                    if scanned:
                        break
                except Exception:
                    _log.debug("excel_unrecognized reader %s failed", reader.__name__, exc_info=True)
                    continue
        else:
            text = _decode_text(file_bytes)
            delimiter = ";" if text.count(";") > text.count(",") else ","
            for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
                if i >= max_scan:
                    break
                scanned.append(row)
    except Exception:
        _log.warning("read_tabular_preview failed for %s", filename, exc_info=True)
        return []
    return scanned


def detect_header_row(
    file_bytes: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    max_scan: int = 15,
) -> int | None:
    """Guess the 1-based row that contains Turkish/English statement column titles.

    Many TR bank exports put 'Ekstre' / 'Statement' on row 1 and 'Tarih' on row 2 or 3.
    """
    header_markers = (
        "tarih",
        "date",
        "aciklama",
        "borc",
        "alacak",
        "tutar",
        "bakiye",
    )
    scanned_rows = _scan_rows_for_header(
        file_bytes, filename, sheet_name=sheet_name, max_scan=max_scan
    )
    if not scanned_rows:
        return None

    best_row: int | None = None
    best_score = 0
    for idx, row in enumerate(scanned_rows):
        cells = [_header_lower(c) for c in row if str(c).strip()]
        joined = " ".join(cells)
        score = sum(1 for m in header_markers if m in joined)
        if "tarih" in joined or "date" in joined:
            score += 2
        if score > best_score:
            best_score = score
            best_row = idx + 1
    return best_row if best_score >= 2 else None


def list_excel_sheets(file_bytes: bytes) -> list[str]:
    """Return sheet names from a real .xlsx workbook (ZIP-based)."""
    if not _is_xlsx_bytes(file_bytes):
        return []
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _flatten_column_name(col: Any) -> str:
    """Flatten MultiIndex / merged Excel headers like ('Statement', 'Tarih')."""
    if isinstance(col, tuple):
        parts = [
            str(p).strip()
            for p in col
            if p is not None
            and str(p).strip()
            and str(p).strip().lower() not in ("nan", "none")
        ]
        if not parts:
            return "Unnamed"
        # Prefer the leaf name (e.g. 'Tarih' under 'Statement')
        for p in reversed(parts):
            low = p.lower()
            if any(kw in low for kw in ("tarih", "date", "açıklama", "aciklama", "borç", "borc", "alacak", "tutar", "bakiye")):
                return p
        return parts[-1] if len(parts) == 1 else " ".join(parts)
    return str(col).strip().replace("\n", " ")


def _normalize_headers(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [_flatten_column_name(c) for c in df.columns]
    df.columns = [" ".join(str(c).split()) for c in df.columns]
    return df


def _read_xlsx_dataframe(
    file_bytes: bytes,
    *,
    header_row: int,
    sheet_name: str | None,
) -> pd.DataFrame:
    header_idx = max(0, header_row - 1)
    bio = io.BytesIO(file_bytes)
    kwargs: dict[str, Any] = {"header": header_idx, "engine": "openpyxl"}
    if sheet_name is not None and sheet_name != "":
        kwargs["sheet_name"] = sheet_name
    try:
        df = pd.read_excel(bio, **kwargs)
    except Exception as exc:
        msg = str(exc).lower()
        if "zip" in msg or "not a zip" in msg:
            raise ValueError(
                "Could not read this Excel file (invalid or wrong format). "
                "Save as .xlsx (Excel 2007+) or upload CSV instead."
            ) from exc
        raise ValueError(f"Could not read Excel file: {exc}") from exc
    return _normalize_headers(df)


def _read_html_dataframe(file_bytes: bytes, *, header_row: int) -> pd.DataFrame:
    df = _pick_largest_table(_read_html_tables_raw(file_bytes))
    return _normalize_headers(_apply_header_row(df, header_row))


def _read_spreadsheetml_dataframe(file_bytes: bytes, *, header_row: int) -> pd.DataFrame:
    df = _read_spreadsheetml_raw(file_bytes)
    return _normalize_headers(_apply_header_row(df, header_row))


def _read_dataframe(
    file_bytes: bytes,
    filename: str,
    *,
    header_row: int = 1,
    sheet_name: str | None = None,
) -> pd.DataFrame:
    fmt = detect_file_format(file_bytes, filename)
    if fmt == "xlsx":
        return _read_xlsx_dataframe(file_bytes, header_row=header_row, sheet_name=sheet_name)
    if fmt == "xls_ole":
        raise ValueError(
            "Legacy .xls (OLE) bank export detected. Open in Excel and Save As .xlsx, or export CSV."
        )
    if fmt == "html":
        return _read_html_dataframe(file_bytes, header_row=header_row)
    if fmt == "spreadsheetml":
        return _read_spreadsheetml_dataframe(file_bytes, header_row=header_row)
    if fmt == "excel_unrecognized":
        errors: list[str] = []
        for label, reader in (
            ("HTML", _read_html_dataframe),
            ("SpreadsheetML", _read_spreadsheetml_dataframe),
        ):
            try:
                return reader(file_bytes, header_row=header_row)
            except Exception as exc:
                errors.append(f"{label}: {exc}")
        raise ValueError(
            "File has an Excel extension but is not a real .xlsx workbook (common with Turkish bank exports). "
            "Open it in Excel and use File → Save As → Excel Workbook (.xlsx), or export CSV. "
            f"Tried: {'; '.join(errors)}"
        )
    header_idx = max(0, header_row - 1)
    text = _decode_text(file_bytes)
    delimiter = ";" if text.count(";") > text.count(",") else ","
    return _normalize_headers(pd.read_csv(io.StringIO(text), header=header_idx, sep=delimiter))


def _parse_date(raw: Any) -> datetime.date | None:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    if isinstance(raw, datetime.datetime):
        return raw.date()
    if isinstance(raw, datetime.date):
        return raw
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        try:
            ts = pd.to_datetime(float(raw), unit="D", origin="1899-12-30", errors="coerce")
            if not pd.isna(ts):
                return ts.date()
        except (ValueError, TypeError, OverflowError):
            pass
    if hasattr(raw, "to_pydatetime"):
        try:
            return raw.to_pydatetime().date()
        except (ValueError, TypeError, AttributeError):
            pass
    s = str(raw).strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        parsed = pd.to_datetime(s, dayfirst=True, errors="coerce")
        if pd.isna(parsed):
            return None
        return parsed.date()
    except (ValueError, TypeError):
        return None


def _cell_raw(row: pd.Series, col: str | None) -> Any:
    if not col or col not in row.index:
        return None
    val = row[col]
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    return val


def _parse_signed_amount(raw: Any) -> float | None:
    """Parse a single amount cell that may be negative (withdrawal)."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return round(float(raw), 2)
    s = str(raw).strip().replace("\xa0", "")
    if not s:
        return None
    negative = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    val = parse_amount_str(s)
    if val is None:
        return None
    if negative or s.startswith("-"):
        return -abs(val)
    return val


def _cell_str(row: pd.Series, col: str | None) -> str:
    val = _cell_raw(row, col)
    if val is None:
        return ""
    return str(val).strip()


def _resolve_debit_credit(
    row: pd.Series, column_mapping: dict[str, str | None]
) -> tuple[float | None, float | None]:
    """Resolve debit/credit from separate columns or one signed Tutar / İşlem Tutarı column."""
    amount_col = column_mapping.get("amount")
    debit_col = column_mapping.get("debit")
    credit_col = column_mapping.get("credit")
    amount_cols = {c for c in (amount_col, debit_col, credit_col) if c}
    single_signed = len(amount_cols) == 1 or (
        debit_col and credit_col and debit_col == credit_col
    )

    if single_signed:
        col = amount_col or debit_col or credit_col
        signed = _parse_signed_amount(_cell_raw(row, col))
        if signed is None or signed == 0:
            return None, None
        if signed < 0:
            return abs(signed), None
        return None, signed

    debit = parse_amount_str(_cell_str(row, debit_col))
    credit = parse_amount_str(_cell_str(row, credit_col))

    if amount_col and not debit and not credit:
        signed = _parse_signed_amount(_cell_raw(row, amount_col))
        if signed is not None:
            if signed < 0:
                debit = abs(signed)
            elif signed > 0:
                credit = signed

    if debit is not None and debit < 0:
        return abs(debit), None
    if credit is not None and credit < 0:
        return abs(credit), None
    if debit == 0:
        debit = None
    if credit == 0:
        credit = None
    return debit, credit


def _raw_line_from_row(row: pd.Series) -> str:
    parts = [str(v) if not (isinstance(v, float) and pd.isna(v)) else "" for v in row.tolist()]
    return delimiter_join(parts)


def delimiter_join(parts: list[str]) -> str:
    buf = io.StringIO()
    csv.writer(buf).writerow(parts)
    return buf.getvalue().strip()


def parse_bank_statement(
    file_bytes: bytes,
    filename: str,
    column_mapping: dict[str, str | None],
    *,
    currency: str = "TRY",
    header_row: int = 1,
    sheet_name: str | None = None,
) -> list[dict]:
    """Parse all rows into dicts ready for persistence."""
    df = _read_dataframe(
        file_bytes, filename, header_row=header_row, sheet_name=sheet_name
    )
    rows: list[dict] = []
    for idx in range(len(df)):
        row = df.iloc[idx]
        import_row_index = idx + 1
        desc = _cell_str(row, column_mapping.get("description"))
        debit, credit = _resolve_debit_credit(row, column_mapping)
        balance = parse_amount_str(_cell_str(row, column_mapping.get("balance")))
        bank_ref = _cell_str(row, column_mapping.get("bank_reference")) or None
        row_date = _parse_date(_cell_raw(row, column_mapping.get("date")))

        amount = 0.0
        original = 0.0
        if debit and debit > 0:
            amount = round(debit, 2)
            original = amount
        elif credit and credit > 0:
            amount = round(credit, 2)
            original = amount

        parsed_ok = row_date is not None and amount > 0 and bool(desc or bank_ref)
        parse_error = None
        if not parsed_ok:
            errs = []
            if row_date is None:
                errs.append("invalid_date")
            if amount <= 0:
                errs.append("invalid_amount")
            if not desc and not bank_ref:
                errs.append("missing_description")
            parse_error = ",".join(errs)

        norm_desc = normalize_description(desc)
        rows.append(
            {
                "import_row_index": import_row_index,
                "date": row_date,
                "description": desc or bank_ref or "",
                "debit_amount": debit if debit else None,
                "credit_amount": credit if credit else None,
                "amount": amount,
                "balance_after": balance,
                "currency": currency,
                "original_amount": original,
                "bank_reference": bank_ref,
                "raw_line_text": _raw_line_from_row(row),
                "normalized_description": norm_desc,
                "parsed_successfully": parsed_ok,
                "parse_error": parse_error,
                "status": "parse_error" if not parsed_ok else "staging",
            }
        )
    return rows


def mapping_to_json(column_mapping: dict[str, str | None]) -> str:
    return json.dumps({k: v for k, v in column_mapping.items() if v})
